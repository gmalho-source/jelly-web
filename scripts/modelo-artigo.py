"""
O modelo Word dos artigos do blog: public/modelos/modelo-artigo-jelly.docx.

É a norma posta em documento. Quem escreve parte dele, troca os textos de
exemplo pelos seus, preenche a ficha técnica do fim, e a importação do painel
(«Importar Markdown ou Word») faz o resto: o estilo «Título» dá o título e o
slug, os «Título 1» e «Título 2» dão as secções, a ficha dá a data, o autor, a
categoria, as etiquetas, a capa e o resumo.

Os valores da ficha estão entre parênteses rectos de propósito: a importação
trata «[...]» como vazio, e um modelo esquecido não cria um autor chamado
«[Nome do autor]». Ver src/lib/ficha-tecnica.ts.

    python3 scripts/modelo-artigo.py     (precisa do python-docx)
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

VERMELHO = RGBColor(0xDD, 0x36, 0x4A)
TINTA = RGBColor(0x15, 0x17, 0x19)
CINZA = RGBColor(0x8A, 0x93, 0xA0)

CATEGORIAS = (
    "Audiovisuais, Branding, Consultoria, Eventos, Inteligência Artificial, Jelly, Marketing, "
    "Opinião, Paid Media, Proteção de Dados, Sala de Imprensa, Social Media, Tecnologia"
)

d = Document()

normal = d.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = TINTA
for nome, tamanho in (("Title", 26), ("Heading 1", 16), ("Heading 2", 13)):
    estilo = d.styles[nome]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(tamanho)
    estilo.font.bold = True
    estilo.font.color.rgb = VERMELHO if nome == "Title" else TINTA


def paragrafo(*partes):
    """Um parágrafo com pedaços: texto simples, ("negrito", texto) ou ("italico", texto)."""
    p = d.add_paragraph()
    for parte in partes:
        if isinstance(parte, tuple):
            forma, texto = parte
            run = p.add_run(texto)
            run.bold = forma == "negrito"
            run.italic = forma == "italico"
        else:
            p.add_run(parte)
    return p


def tabela(linhas, larguras=None, cabecalho=True):
    t = d.add_table(rows=len(linhas), cols=len(linhas[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for r, linha in enumerate(linhas):
        for c, texto in enumerate(linha):
            celula = t.cell(r, c)
            celula.text = ""
            run = celula.paragraphs[0].add_run(texto)
            if (cabecalho and r == 0) or (larguras and c == 0):
                run.bold = True
            if texto.startswith("["):
                run.font.color.rgb = CINZA
    return t


# ── O artigo ────────────────────────────────────────────────────────────────

d.add_paragraph("Título do artigo: curto, claro e com a ideia principal", style="Title")

paragrafo(
    "Primeiro parágrafo. É o que abre o artigo no site, com a letra capitular vermelha: diz logo ao que o "
    "texto vem, em duas ou três frases. O título de cima usa o estilo ",
    ("negrito", "«Título»"),
    " do Word — é dele que sai o título do artigo e o endereço (o slug).",
)

d.add_heading("Um título de secção", level=1)
paragrafo(
    "As secções usam o estilo ",
    ("negrito", "«Título 1»"),
    ", e as subsecções o ",
    ("negrito", "«Título 2»"),
    ". Não faças títulos à mão com letra maior ou a negrito: o site só reconhece os estilos.",
)

d.add_heading("Uma subsecção", level=2)
paragrafo(
    "Negrito, ",
    ("italico", "itálico"),
    " e ligações passam para o site como estão. As imagens colam-se no sítio onde devem aparecer, cada uma "
    "no seu parágrafo; o texto alternativo põe-se com o botão direito na imagem, em «Editar texto alternativo».",
)
for texto in (
    "As listas usam os botões de lista do Word.",
    "Com marcas ou numeradas.",
    "Um ponto por linha.",
):
    d.add_paragraph(texto, style="List Bullet")

d.add_heading("Tabelas", level=2)
paragrafo("Uma tabela simples passa para o site como tabela. A primeira linha a negrito é o cabeçalho.")
tabela([["Canal", "2024", "2025"], ["Loja online", "4%", "18%"], ["Lojas próprias", "96%", "82%"]])
paragrafo("O último parágrafo do artigo fica aqui, antes da ficha técnica.")

# ── A ficha técnica ─────────────────────────────────────────────────────────

d.add_heading("Ficha técnica", level=1)
ficha = tabela(
    [
        ["Campo", "Valor"],
        ["Data", "[Data de publicação, por exemplo 28/09/2026]"],
        ["Autor", "[Nome tal como está no painel, em Editorial → Autores]"],
        ["Categoria", f"[Uma só. Por exemplo: {CATEGORIAS}]"],
        ["Etiquetas", "[Separadas por vírgulas: inteligência artificial, automação, marketing]"],
        ["Capa", "[Cola aqui a imagem de capa, horizontal, com 1600 px de largura ou mais]"],
        ["Resumo", "[Opcional. Uma frase até 155 caracteres, a que sai no Google. Vazio, usa-se o botão de IA]"],
        ["Slug", "[Opcional. Vazio, sai do título]"],
        [
            "Nota",
            "Esta ficha sai do artigo ao importar. Não mudes o título «Ficha técnica» nem os nomes da coluna da "
            "esquerda, e apaga os parênteses rectos quando preencheres. O que ficar entre parênteses não é lido.",
        ],
    ],
    larguras=True,
)

destino = Path(__file__).resolve().parent.parent / "public" / "modelos" / "modelo-artigo-jelly.docx"
d.save(destino)
print(destino)
